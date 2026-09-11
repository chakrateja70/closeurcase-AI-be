"""Tests for uploaded-document handling in src/utils/document.py.

Summarization quality itself is checked against live API calls, not here. What
these pin are the decisions made *before* a model is ever called, all of which
are cheap to get wrong and silent when wrong:

- identification by content, so a renamed file cannot pick its own handler,
- the scanned-vs-digital threshold, which decides between a cheap text call and
  an expensive per-page one, and
- the rejections (empty, oversized, encrypted, too long) that exist to stop a
  bad upload from costing money.

PDFs are built by hand rather than loaded from fixtures so the text content -
and therefore the expected extraction - is visible in the test itself.
"""

import io
import zipfile

import docx
import pytest
from pypdf import PdfReader, PdfWriter

from src.utils import document
from src.utils.document import (
    BRANCH_FILE,
    BRANCH_IMAGE,
    BRANCH_TEXT,
    KIND_DOCX,
    KIND_JPEG,
    KIND_PDF,
    KIND_PNG,
    KIND_TEXT,
    DocumentError,
    extract,
    from_text,
    sniff,
)

PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
JPEG_BYTES = b"\xff\xd8\xff\xe0" + b"\x00" * 64


def make_pdf(page_texts: list[str]) -> bytes:
    """A minimal but structurally valid multi-page PDF with a real text layer.

    Object numbering is contiguous (catalog, pages, font, then a page/content
    pair each) so the xref table below can be written by walking the offsets in
    order. A page given "" gets an empty text stream, which is what a scan looks
    like to an extractor: pages present, no text on them.
    """
    page_ids = [4 + 2 * i for i in range(len(page_texts))]
    content_ids = [5 + 2 * i for i in range(len(page_texts))]

    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids ["
        + b" ".join(b"%d 0 R" % pid for pid in page_ids)
        + b"] /Count %d >>" % len(page_texts),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for page_id, content_id, text in zip(page_ids, content_ids, page_texts):
        objects[page_id] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_id
        )
        escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        stream = b"BT /F1 12 Tf 72 720 Td (" + escaped.encode("latin-1") + b") Tj ET"
        objects[content_id] = (
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        )

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += b"%d 0 obj\n" % number + objects[number] + b"\nendobj\n"

    xref_offset = len(out)
    size = max(objects) + 1
    out += b"xref\n0 %d\n0000000000 65535 f \n" % size
    for number in range(1, size):
        out += b"%010d 00000 n \n" % offsets[number]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        size,
        xref_offset,
    )
    return bytes(out)


def make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    buffer = io.BytesIO()
    written = docx.Document()
    for paragraph in paragraphs:
        written.add_paragraph(paragraph)
    if table_rows:
        table = written.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row, values in zip(table.rows, table_rows):
            for cell, value in zip(row.cells, values):
                cell.text = value
    written.save(buffer)
    return buffer.getvalue()


# A page of plausible pleading text, comfortably over MIN_CHARS_PER_PAGE.
PLEADING = (
    "IN THE COURT OF THE CIVIL JUDGE AT VISAKHAPATNAM. Suit No. 412 of 2024. "
    "The plaintiff states that a cheque dated 12.03.2024 for Rs. 4,50,000 was "
    "dishonoured for insufficiency of funds and that a statutory notice was "
    "issued on 20.03.2024 to which no reply has been received."
)


# --- Identification ---------------------------------------------------------


# `ids` on every bytes parametrize below: pytest puts the generated id into an
# environment variable, and a repr of a whole PDF blows past the Windows 32 KB
# limit for one.
@pytest.mark.parametrize(
    "data, expected",
    [
        (make_pdf([PLEADING]), KIND_PDF),
        (make_docx(["hello"]), KIND_DOCX),
        (PNG_BYTES, KIND_PNG),
        (JPEG_BYTES, KIND_JPEG),
    ],
    ids=["pdf", "docx", "png", "jpeg"],
)
def test_sniff_identifies_accepted_types(data, expected):
    assert sniff(data) == expected


@pytest.mark.parametrize(
    "data",
    [
        b"MZ\x90\x00" + b"\x00" * 64,  # a PE executable renamed to .pdf
        b"plain text pretending to be a document",
        b"GIF89a" + b"\x00" * 64,  # an image, but not one we accept
    ],
    ids=["executable", "plain_text", "gif"],
)
def test_sniff_rejects_by_content_not_by_name(data):
    """Nothing here consults a filename or a declared content type - both are
    set by the client, so neither is evidence."""
    with pytest.raises(DocumentError):
        sniff(data)


def test_sniff_rejects_a_zip_that_is_not_a_docx():
    """A .docx is a zip, but so is every other Office file and any archive; the
    word/ entry is the thing that distinguishes it."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
    with pytest.raises(DocumentError):
        sniff(buffer.getvalue())


# --- Rejections -------------------------------------------------------------


def test_empty_upload_is_rejected():
    with pytest.raises(DocumentError):
        extract(b"")


def test_oversized_upload_is_rejected_before_parsing():
    with pytest.raises(DocumentError, match="too large"):
        extract(b"%PDF-" + bytes(document.MAX_FILE_BYTES))


def test_encrypted_pdf_is_reported_not_mistaken_for_a_scan():
    """An encrypted PDF extracts as empty pages, which is indistinguishable
    from a scan - if this check were dropped it would silently take the
    expensive per-page branch instead of telling the user to unlock it."""
    writer = PdfWriter(clone_from=io.BytesIO(make_pdf([PLEADING])))
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(DocumentError, match="password"):
        extract(buffer.getvalue())


def test_pdf_over_the_page_cap_is_rejected():
    data = make_pdf([PLEADING] * (document.MAX_PAGES + 1))
    with pytest.raises(DocumentError, match="pages"):
        extract(data)


def test_docx_with_no_text_is_rejected():
    with pytest.raises(DocumentError, match="No readable text"):
        extract(make_docx([]))


def test_unreadable_pdf_is_rejected():
    with pytest.raises(DocumentError):
        extract(b"%PDF-1.4\nnot actually a pdf")


# --- Branch selection -------------------------------------------------------


def test_digital_pdf_takes_the_text_branch_with_page_markers():
    result = extract(make_pdf([PLEADING, PLEADING]))

    assert result.branch == BRANCH_TEXT
    assert result.page_count == 2
    assert "--- Page 1 ---" in result.text
    assert "--- Page 2 ---" in result.text
    assert "Suit No. 412 of 2024" in result.text
    assert result.truncated is False


def test_pdf_with_no_text_layer_takes_the_file_branch():
    """A scan: pages exist, no text on them. It goes to the model as the file
    itself rather than being rejected or OCR'd."""
    result = extract(make_pdf(["", "", ""]))

    assert result.branch == BRANCH_FILE
    assert result.page_count == 3
    assert result.text is None


def test_threshold_is_averaged_not_per_page():
    """One sparse page - a cover sheet or an index - must not tip a document
    that is otherwise full of text onto the expensive branch."""
    result = extract(make_pdf(["", PLEADING, PLEADING]))

    assert result.branch == BRANCH_TEXT


def test_docx_takes_the_text_branch_including_table_cells():
    """Tables carry schedules of property, dates and payment particulars in
    legal drafting - dropping them would lose the specifics worth extracting."""
    data = make_docx(
        ["Statement of claim."], table_rows=[["Item", "Amount"], ["Rent", "45,000"]]
    )
    result = extract(data)

    assert result.branch == BRANCH_TEXT
    assert "Statement of claim." in result.text
    assert "Rent | 45,000" in result.text


@pytest.mark.parametrize("data", [PNG_BYTES, JPEG_BYTES], ids=["png", "jpeg"])
def test_images_take_the_image_branch(data):
    result = extract(data)

    assert result.branch == BRANCH_IMAGE
    assert result.page_count is None
    assert result.text is None


def test_long_text_is_truncated_and_says_so(monkeypatch):
    monkeypatch.setattr(document, "MAX_TEXT_CHARS", 50)
    result = extract(make_pdf([PLEADING]))

    assert result.truncated is True
    assert len(result.text) == 50


# --- Metadata ---------------------------------------------------------------


def test_media_type_is_carried_for_the_attachment_branches():
    """Providers switch on this to decide how to attach the bytes; base64
    encoding is theirs to do, not this module's."""
    assert extract(PNG_BYTES).media_type == "image/png"
    assert extract(make_pdf([""])).media_type == "application/pdf"


def test_pdf_page_count_matches_the_reader():
    data = make_pdf([PLEADING] * 4)
    assert extract(data).page_count == len(PdfReader(io.BytesIO(data)).pages)


# --- Pasted text ------------------------------------------------------------
#
# The second entry point. It shares every rule `extract` applies except the two
# that have no meaning for characters: sniffing and page counting.


def test_pasted_text_takes_the_text_branch():
    result = from_text(PLEADING)

    assert result.branch == BRANCH_TEXT
    assert result.kind == KIND_TEXT
    assert result.text == PLEADING
    assert result.truncated is False


def test_pasted_text_has_no_pages():
    """Not an omission. Pasted text genuinely has no pages, so every entry ends
    up `page_status: unavailable` - inventing markers to fill `source_pages`
    would publish page numbers that refer to nothing."""
    result = from_text(PLEADING)

    assert result.page_count is None
    assert "--- Page" not in result.text


def test_pasted_text_is_stripped():
    result = from_text(f"\n\n  {PLEADING}  \n\n")
    assert result.text == PLEADING


@pytest.mark.parametrize("value", ["", "   ", "\n\t "])
def test_empty_text_is_rejected(value):
    with pytest.raises(document.DocumentError, match="No text"):
        from_text(value)


def test_text_below_the_floor_is_rejected():
    """A file at least had to be a real PDF to get this far; a text box will
    happily submit 'hi'. The call would be billed either way."""
    with pytest.raises(document.DocumentError, match="too short"):
        from_text("x" * (document.MIN_TEXT_CHARS - 1))


def test_text_at_the_floor_is_accepted():
    assert from_text("x" * document.MIN_TEXT_CHARS).text


def test_long_pasted_text_is_truncated_and_says_so(monkeypatch):
    """The same cap a file is held to - a paste is not a way around it."""
    monkeypatch.setattr(document, "MAX_TEXT_CHARS", 250)
    result = from_text("x" * 400)

    assert result.truncated is True
    assert len(result.text) == 250


def test_pasted_text_carries_its_bytes_for_the_size_log():
    """Nothing sends them - the TEXT branch travels as `text` - but the log
    line reports size_bytes for every document, whichever way it arrived."""
    result = from_text(PLEADING)

    assert result.data == PLEADING.encode("utf-8")
    assert result.size_bytes == len(PLEADING.encode("utf-8"))
