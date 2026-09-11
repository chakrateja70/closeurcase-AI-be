"""Uploaded-document handling: identify, validate, and get text out.

Generic and domain-free, like `helper.py` - nothing here knows what a petition
is. It answers one question: given raw uploaded bytes, what is this file and
how should it be put in front of a model?

    bytes -> sniff -> PDF   -> pypdf text -> enough text? -> TEXT branch
                   |                      \\-> too little  -> FILE branch
                   |-> DOCX  -> python-docx --------------> TEXT branch
                   \\-> JPEG/PNG -----------------------------> IMAGE branch

    str ------------------------------------------------------> TEXT branch

`from_text` is the second entry point, for a document pasted into a box rather
than uploaded. It skips sniffing (there is nothing to identify) and page
counting (pasted text has no pages), but shares every other rule: the same
length cap, the same TEXT branch, the same `ExtractedDocument`. Downstream code
therefore cannot tell the two apart, which is the point - a paste is a document
that happens to have arrived as characters.

The branch tells the caller how to attach the document to a model request;
`summary_providers` translates it into whatever the chosen SDK wants. Nothing
here knows about any particular provider - that is why the base64 encoding this
module used to do now lives with the provider that needs it.

Why the FILE branch exists: court filings are very often scans or photocopies
of a signed, stamped original, and a scanned PDF has no text layer at all -
extraction returns empty pages. Rather than run OCR infrastructure, those go to
the model as the file itself, which reads the rendered pages directly.

PDF text extraction is `pypdf`, which is BSD-licensed. Do NOT replace it with
PyMuPDF/`fitz`: that is AGPL, and linking it into a hosted commercial service
carries a source-disclosure obligation. `pdfplumber` (MIT) is the safe upgrade
if layout-aware extraction is ever needed.

Failures raise `DocumentError`, a plain ValueError subclass. This module stays
free of FastAPI so it can be tested and reused without a request; translating
to an HTTP status is the calling service's job.
"""

from __future__ import annotations

import io
import logging
import zipfile
from dataclasses import dataclass

import docx
from pypdf import PdfReader

logger = logging.getLogger(__name__)

# --- Limits -----------------------------------------------------------------
#
# These bound the DOCUMENT, not the model call (timeouts and retries live with
# the service that makes it). Every one of them is a cost ceiling as much as a
# safety one: on the FILE branch the provider bills per rendered page.

MAX_FILE_BYTES = 20 * 1024 * 1024  # 20 MB
MAX_PAGES = 30
MAX_TEXT_CHARS = 250_000

# Average extracted characters per page below which a PDF is treated as a scan.
# A genuinely sparse text page (a cover sheet, an index) still clears this once
# averaged over the document; a scan yields ~0.
MIN_CHARS_PER_PAGE = 100

# Branches: how the document should be attached to a model request.
BRANCH_TEXT = "text"
BRANCH_FILE = "file"
BRANCH_IMAGE = "image"

KIND_PDF = "pdf"
KIND_DOCX = "docx"
KIND_JPEG = "jpeg"
KIND_PNG = "png"
KIND_TEXT = "text"

MEDIA_TYPES = {
    KIND_PDF: "application/pdf",
    KIND_DOCX: (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
    KIND_JPEG: "image/jpeg",
    KIND_PNG: "image/png",
    KIND_TEXT: "text/plain",
}

ACCEPTED_DESCRIPTION = "PDF, DOCX, JPG or PNG"

# Pasted text is capped by the same MAX_TEXT_CHARS as an extracted file, but it
# also needs a floor: a file at least had to be a real PDF to get this far,
# whereas a text box will happily submit "hi". Below this there is nothing to
# summarise and the call would be paid for anyway.
MIN_TEXT_CHARS = 200


class DocumentError(ValueError):
    """The upload cannot be processed. The message is caller-facing, so it must
    stay free of internals - it is shown to whoever sent the file."""


@dataclass(frozen=True)
class ExtractedDocument:
    """What one upload became.

    `text` is set on the TEXT branch only; the other two branches send the raw
    bytes, so the caller reads `data` and `media_type` instead. `page_count` is
    None for anything that has no pages (DOCX reflows, images are one image).
    """

    kind: str
    branch: str
    media_type: str
    data: bytes
    text: str | None = None
    page_count: int | None = None
    truncated: bool = False

    @property
    def size_bytes(self) -> int:
        return len(self.data)


# --- Identification ---------------------------------------------------------
#
# By content, never by the declared Content-Type or the filename extension:
# both are set by the client and neither is evidence of anything.

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"
_JPEG_MAGIC = b"\xff\xd8\xff"
_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


def sniff(data: bytes) -> str:
    """Detected kind constant, or raise `DocumentError` for anything else."""
    if data.startswith(_PDF_MAGIC):
        return KIND_PDF
    if data.startswith(_JPEG_MAGIC):
        return KIND_JPEG
    if data.startswith(_PNG_MAGIC):
        return KIND_PNG
    if data.startswith(_ZIP_MAGIC) and _is_docx_zip(data):
        return KIND_DOCX
    raise DocumentError(
        f"Unsupported file type. Please upload a {ACCEPTED_DESCRIPTION} file."
    )


def _is_docx_zip(data: bytes) -> bool:
    """A .docx is a zip; so is .xlsx, .pptx, .jar and an ordinary archive. The
    word/ entry is what distinguishes it."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return "word/document.xml" in archive.namelist()
    except (zipfile.BadZipFile, OSError):
        return False


# --- Extraction -------------------------------------------------------------


def extract(data: bytes) -> ExtractedDocument:
    """Validate an upload and turn it into the branch that should be sent.

    Raises `DocumentError` for an empty, oversized, unsupported, encrypted, too
    long, or unreadable file - every rejection happens here, before anything
    reaches a model and costs money.
    """
    if not data:
        raise DocumentError("The uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError(
            f"File is too large ({len(data) / 1_048_576:.1f} MB). "
            f"The limit is {MAX_FILE_BYTES // 1_048_576} MB."
        )

    kind = sniff(data)
    media_type = MEDIA_TYPES[kind]

    if kind in (KIND_JPEG, KIND_PNG):
        return ExtractedDocument(
            kind=kind, branch=BRANCH_IMAGE, media_type=media_type, data=data
        )

    if kind == KIND_DOCX:
        text, truncated = _truncate(_extract_docx_text(data))
        if not text:
            raise DocumentError(
                "No readable text was found in this document. If it is a scan, "
                "please upload it as a PDF or an image instead."
            )
        return ExtractedDocument(
            kind=kind,
            branch=BRANCH_TEXT,
            media_type=media_type,
            data=data,
            text=text,
            truncated=truncated,
        )

    return _extract_pdf(data, media_type)


def from_text(text: str) -> ExtractedDocument:
    """Turn pasted text into the same `ExtractedDocument` an upload becomes.

    Every rejection happens here, as it does in `extract`, before anything
    reaches a model: too short to be a document, or past the same length cap a
    file is held to.

    `page_count` stays None and no `--- Page N ---` markers are added. That is
    not an omission - pasted text genuinely has no pages, so every entry comes
    back `page_status: unavailable`, exactly as a DOCX does. Inventing markers
    to fill the field would publish page numbers that refer to nothing.

    `data` carries the UTF-8 bytes so `size_bytes` still means something in the
    log line; nothing sends them, because the TEXT branch travels as `text`.
    """
    cleaned = text.strip()
    if not cleaned:
        raise DocumentError("No text was provided.")
    if len(cleaned) < MIN_TEXT_CHARS:
        raise DocumentError(
            f"This text is too short to summarise ({len(cleaned)} characters). "
            f"Please paste at least {MIN_TEXT_CHARS} characters, or upload the "
            f"document itself."
        )

    capped, truncated = _truncate(cleaned)
    return ExtractedDocument(
        kind=KIND_TEXT,
        branch=BRANCH_TEXT,
        media_type=MEDIA_TYPES[KIND_TEXT],
        data=capped.encode("utf-8"),
        text=capped,
        truncated=truncated,
    )


def _extract_pdf(data: bytes, media_type: str) -> ExtractedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        encrypted = reader.is_encrypted
    except Exception as exc:  # pypdf raises a family of read/parse errors
        raise DocumentError(
            "This PDF could not be read. It may be corrupt or incomplete."
        ) from exc

    # Must be checked before anything touches the pages: on an encrypted file
    # `reader.pages` itself raises, so a later check would report it as corrupt.
    # And were the check dropped entirely, a decryptable-but-locked file would
    # extract as empty pages - indistinguishable from a scan - and get silently
    # routed to the far more expensive file branch instead of being reported.
    if encrypted:
        raise DocumentError(
            "This PDF is password protected. Please upload an unlocked copy."
        )

    try:
        page_count = len(reader.pages)
    except Exception as exc:
        raise DocumentError(
            "This PDF could not be read. It may be corrupt or incomplete."
        ) from exc

    if page_count == 0:
        raise DocumentError("This PDF has no pages.")
    if page_count > MAX_PAGES:
        raise DocumentError(
            f"This document has {page_count} pages. "
            f"The limit is {MAX_PAGES} pages per upload."
        )

    pages = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:  # one broken page should not lose the other 29
            logger.warning("pdf: page %d could not be extracted", index)
            page_text = ""
        pages.append((index, page_text.strip()))

    extracted = sum(len(text) for _, text in pages)
    if extracted / page_count < MIN_CHARS_PER_PAGE:
        # No usable text layer: a scan or a photocopy. Send the file itself and
        # let the model read the rendered pages.
        return ExtractedDocument(
            kind=KIND_PDF,
            branch=BRANCH_FILE,
            media_type=media_type,
            data=data,
            page_count=page_count,
        )

    text, truncated = _truncate(_join_pages(pages))
    return ExtractedDocument(
        kind=KIND_PDF,
        branch=BRANCH_TEXT,
        media_type=media_type,
        data=data,
        text=text,
        page_count=page_count,
        truncated=truncated,
    )


def _join_pages(pages: list[tuple[int, str]]) -> str:
    """Explicit page markers so a model asked to cite page numbers has
    something real to cite; without them it can only guess."""
    return "\n\n".join(
        f"--- Page {number} ---\n{text}" for number, text in pages if text
    )


def _extract_docx_text(data: bytes) -> str:
    """Paragraphs plus table cells. Tables are not decoration in legal drafting
    - schedules of property, lists of dates and payment particulars all live in
    them, so dropping them would lose the specifics worth extracting."""
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(
            "This Word document could not be read. It may be corrupt."
        ) from exc

    blocks = [p.text.strip() for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(block for block in blocks if block)


def _truncate(text: str) -> tuple[str, bool]:
    """Cap the text branch. Returns the text and whether anything was dropped,
    so the caller can say so rather than silently summarising a partial file."""
    if len(text) <= MAX_TEXT_CHARS:
        return text, False
    return text[:MAX_TEXT_CHARS], True
